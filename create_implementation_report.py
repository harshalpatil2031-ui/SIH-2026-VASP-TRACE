from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

OUT = "VASP_TRACE_Implementation_Report.docx"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tc_pr.append(shd)

def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for side in ('top','left','bottom','right'):
        el = OxmlElement('w:' + side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); b.append(el)
    tc_pr.append(b)

def set_cell(cell, text, bold=False, color=None):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(3); p.paragraph_format.space_before = Pt(3)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(9)
    if color: r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    borders(cell)

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text); r.font.color.rgb = RGBColor(0,0,0)
    return p

def para(doc, text, bold_lead=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.12
    if bold_lead:
        r=p.add_run(bold_lead); r.bold=True
    p.add_run(text)
    return p

def bullet(doc, text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); p.add_run(text); return p

doc = Document()
sec = doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); styles['Normal'].font.size=Pt(10)
for s in ('Title','Heading 1','Heading 2'):
    styles[s].font.name='Aptos'; styles[s]._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); styles[s].font.color.rgb=RGBColor(0,0,0)
styles['Heading 1'].font.size=Pt(15); styles['Heading 2'].font.size=Pt(11)

p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('VASP TRACE Implementation Report')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Live Blockchain Integration and Evidence Management'); r.bold=True; r.font.size=Pt(12)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Prepared for SIH Prototype SIH26182').italic=True

heading(doc,'Executive Summary')
para(doc,'VASP TRACE has been extended from a synthetic demonstration platform into a dual-mode investigation system. The existing demonstration engine remains available for offline presentation and testing, while a separate live path now supports read-only blockchain retrieval, concurrent source enrichment, evidence persistence, integrity verification, and structured export.')
para(doc,'The completed work provides the software foundation for real API-backed tracing. Live operation still requires valid provider credentials and appropriate institutional authorization before it can be used with real cases.')

heading(doc,'Implemented Architecture')
para(doc,'The implementation follows a backend-first design. The browser sends a single trace request to the VASP TRACE backend; API credentials are never exposed to the frontend. The backend selects the operating mode, calls sources, normalizes transactions, constructs the existing investigation graph, calculates attribution, records evidence, and returns the result to the dashboard.')

t=doc.add_table(rows=1, cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
for i,h in enumerate(['Layer','Implemented responsibility','Status']): shade(t.rows[0].cells[i],'17365D'); set_cell(t.rows[0].cells[i],h,True,(255,255,255))
rows=[
('Frontend workbench','Trace-mode choice, displayed run ID, integrity check, evidence download','Complete'),
('Trace API','Single backend entry point for DEMO, LIVE, and AUTO modes','Complete'),
('Provider adapters','Read-only Etherscan, TronGrid, and Bitquery integration boundaries','Complete'),
('Trace coordinator','Bounded multi-hop retrieval and concurrent provider work','Complete'),
('Graph and attribution','Existing NetworkX graph and explainable VASP scoring reused','Complete'),
('Evidence store','SQLite investigation-run and transaction audit records','Complete'),
]
for ri,row in enumerate(rows):
    cells=t.add_row().cells
    for i,v in enumerate(row):
        if ri%2==1: shade(cells[i],'F2F6FA')
        set_cell(cells[i],v, i==2)

heading(doc,'Operating Modes')
para(doc,'The application now makes the source of a trace explicit. This is important because demonstration records must never be confused with real blockchain evidence.')
t=doc.add_table(rows=1, cols=3); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,h in enumerate(['Mode','Behaviour','Evidence meaning']): shade(t.rows[0].cells[i],'17365D'); set_cell(t.rows[0].cells[i],h,True,(255,255,255))
for ri,row in enumerate([
('DEMO','Uses the existing seeded and generated synthetic ledger.','Presentation and testing only.'),
('LIVE','Uses only configured live providers. Failure is reported; no synthetic fallback occurs.','Technical record of retrieved live chain data.'),
('AUTO','Attempts live tracing when chosen and configured.','Mode is returned and stored with the run.'),
]):
    cells=t.add_row().cells
    for i,v in enumerate(row):
        if ri%2: shade(cells[i],'F2F6FA')
        set_cell(cells[i],v)

heading(doc,'Live API Integration')
heading(doc,'Provider Responsibilities',2)
bullet(doc,'Etherscan API V2 adapter retrieves outbound Ethereum ERC-20 USDT transfers and normalizes transaction hash, sender, receiver, amount, timestamp, and block number.')
bullet(doc,'TronGrid V1 adapter retrieves confirmed outbound TRON TRC-20 USDT transfers using the configured API key.')
bullet(doc,'Bitquery GraphQL adapter provides optional multi-chain enrichment for TRON and Ethereum. It runs alongside the direct chain provider when configured.')
bullet(doc,'The direct chain provider is treated as the canonical transfer source. Bitquery supplements the trace and does not by itself create a verified VASP attribution.')
heading(doc,'Concurrent Trace Processing',2)
para(doc,'For every wallet at each permitted hop, the coordinator submits the direct provider request and the Bitquery request concurrently where Bitquery is configured. It uses bounded hop depth, bounded transfer counts, and timeouts. Provider success or failure is captured in the trace provenance rather than hidden from the investigator.')

heading(doc,'Evidence and Audit Implementation')
para(doc,'Every completed trace creates a unique investigation-run record. The database persists the technical inputs required for later review rather than relying on the changing state of an external API.')
t=doc.add_table(rows=1, cols=2); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,h in enumerate(['Stored item','Purpose']): shade(t.rows[0].cells[i],'17365D'); set_cell(t.rows[0].cells[i],h,True,(255,255,255))
for ri,row in enumerate([
('Investigation run ID','Stable reference for one trace execution.'),
('Trace mode and provenance','Distinguishes LIVE from DEMO and records provider outcomes.'),
('Normalized transactions','Preserves the traced wallet-to-wallet transfers.'),
('Provider response hashes','Supports later integrity review of live-source payloads.'),
('Attribution snapshot','Preserves the VASP candidate and factor result at trace time.'),
('SHA-256 manifest','Detects change to the canonical technical trace record.'),
]):
    cells=t.add_row().cells
    for i,v in enumerate(row):
        if ri%2: shade(cells[i],'F2F6FA')
        set_cell(cells[i],v)

heading(doc,'Evidence Package Workflow')
para(doc,'The dashboard now displays the saved investigation run ID after tracing. Investigators can select Run Check to recompute the stored manifest and compare it to the stored SHA-256 value. Selecting Evidence downloads a JSON evidence package named with the run ID.')
para(doc,'The exported package contains the run metadata, provenance, transactions, attribution snapshot, manifest, and a technical-evidence disclaimer. It is intended for internal review, case documentation, and attachment to a lawful VASP or SAHYOG draft. It does not itself establish customer identity, freeze assets, or create court admissibility without authorized handling and attestation.')

heading(doc,'API Endpoints Added')
t=doc.add_table(rows=1, cols=2); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,h in enumerate(['Endpoint','Purpose']): shade(t.rows[0].cells[i],'17365D'); set_cell(t.rows[0].cells[i],h,True,(255,255,255))
for ri,row in enumerate([
('GET /api/sources/status','Shows whether each live source is configured without exposing credentials.'),
('POST /api/trace','Runs DEMO, LIVE, or AUTO trace processing and creates an evidence run.'),
('GET /api/evidence/runs/{run_id}','Retrieves a saved evidence run.'),
('POST /api/evidence/runs/{run_id}/verify','Independently verifies the saved manifest.'),
('GET /api/evidence/runs/{run_id}/export','Exports the structured JSON evidence package.'),
]):
    cells=t.add_row().cells
    for i,v in enumerate(row):
        if ri%2: shade(cells[i],'F2F6FA')
        set_cell(cells[i],v)

heading(doc,'Verification Performed')
bullet(doc,'Python backend compilation completed successfully after implementation.')
bullet(doc,'Existing DEMO trace executed successfully and continued to return four transfers for Case 147.')
bullet(doc,'A trace run was stored in the evidence database with its associated transaction records.')
bullet(doc,'Saved-run retrieval, manifest verification, and evidence-package export were exercised successfully.')

heading(doc,'Configuration Required for Live Operation')
para(doc,'Live operation is intentionally inactive until credentials are supplied locally. The repository includes an .env.example file with the required variable names: ETHERSCAN_API_KEY, TRONGRID_API_KEY, BITQUERY_API_KEY, VASP_TRACE_MODE, TRACE_REQUEST_TIMEOUT_SECONDS, TRACE_MAX_HOPS, and TRACE_MAX_TRANSFERS_PER_WALLET.')

heading(doc,'Current Limits and Recommended Next Steps')
bullet(doc,'Add valid API credentials and test known public wallets in LIVE mode before any case deployment.')
bullet(doc,'Replace local SQLite with secured PostgreSQL for a multi-user or production deployment.')
bullet(doc,'Add authenticated investigator accounts, roles, review approval, and comprehensive access logs.')
bullet(doc,'Add a curated and versioned VASP label registry with documented sources and human review.')
bullet(doc,'Keep SAHYOG as a reviewed draft workflow until official authorized integration access is obtained.')
bullet(doc,'Conduct security, privacy, legal, and operational validation before using real victim or FIR data.')

heading(doc,'Conclusion')
para(doc,'The SIH prototype now has a working architecture for safe migration from a visual synthetic demo to an auditable API-backed forensic workflow. The demonstration experience is preserved, while live tracing, provenance capture, evidence integrity, and investigator export are now represented as explicit, testable components.')

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; fr=footer.add_run('VASP TRACE Implementation Report | SIH26182'); fr.font.size=Pt(8); fr.font.color.rgb=RGBColor(90,90,90)
doc.save(OUT)
